# This module supplies the Word layout for the statistics options report.
# It creates a document with shared page geometry, typography, and table borders.
# Small helpers add prose, headings, explicit pages, and comparison tables.
# The adjacent report builder owns the research content and final save operation.
# Outputs use this module directory so the build remains inside the repository.
# Rendering and visual inspection are separate steps before delivery.
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
s.header_distance = s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
h = s.header.paragraphs[0]
h.text = 'DIGITALCORPORA   •   AWS STATISTICS OPTIONS   •   9 OCTOBER 2026'
h.runs[0].font.size = Pt(8)
f = s.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Planning report  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); f._p.append(fld)
for r in f.runs: r.font.size = Pt(8)

def p(text, bold_lead=False):
    para = doc.add_paragraph()
    if bold_lead and ': ' in text:
        a, b = text.split(': ', 1); para.add_run(a + ': ').bold = True; para.add_run(b)
    else: para.add_run(text)
    return para

def title(text): doc.add_heading(text, 1)
def sub(text): doc.add_heading(text, 2)
def page(text):
    heading = doc.add_heading(text, 1)
    heading.paragraph_format.page_break_before = True

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for cell, value in zip(t.rows[0].cells, headers): cell.text = value
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row): cell.text = str(value)
    for i, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for j, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '323E48' if i == 0 else ('F4F6F8' if i % 2 == 0 else 'FFFFFF')); tc.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            tc.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '80'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tc.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2); para.paragraph_format.line_spacing = 1.02
                if j > 0: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size = Pt(9.5)
                    if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
