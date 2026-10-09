# This file builds the requested AWS statistics options report as a Word document.
# It contains research prose and comparison tables rather than application code.
# The report separates current observations, incident history, and cost assumptions.
# Each section starts on a deliberate page and uses shared typography and table rules.
# References are linked to first-party AWS documentation and public project sources.
# Rendering and visual inspection are performed separately before the document is delivered.
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
s.header_distance = s.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
h = s.header.paragraphs[0]
h.text = 'DIGITALCORPORA   •   AWS STATISTICS OPTIONS   •   9 OCTOBER 2026'
h.runs[0].font.size = Pt(8)
f = s.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Planning report  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); f._p.append(fld)
for r in f.runs: r.font.size = Pt(8)

def p(text, bold_lead=False):
    para = doc.add_paragraph()
    if bold_lead and ': ' in text:
        a, b = text.split(': ', 1); para.add_run(a + ': ').bold = True; para.add_run(b)
    else: para.add_run(text)
    return para

def title(text): doc.add_heading(text, 1)
def sub(text): doc.add_heading(text, 2)
def page(text):
    heading = doc.add_heading(text, 1)
    heading.paragraph_format.page_break_before = True

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for cell, value in zip(t.rows[0].cells, headers): cell.text = value
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(t.add_row().cells, row): cell.text = str(value)
    for i, row in enumerate(t.rows):
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for j, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '323E48' if i == 0 else ('F4F6F8' if i % 2 == 0 else 'FFFFFF')); tc.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            tc.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '80'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tc.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2); para.paragraph_format.line_spacing = 1.02
                if j > 0: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size = Pt(9.5)
                    if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

doc.add_paragraph('DigitalCorpora statistics system options', 'Title')
p('Prepared for Simson Garfinkel by Codex  •  9 October 2026')
p('Recommendation: preserve the free S3 download route and replace the per-request MySQL statistics store with daily AWS aggregates. Start with S3 Storage Lens Advanced and a small scheduled publisher. CloudFront would add paid delivery; consider it only if its bot visibility is worth that expense.', True)
p('AWS can already provide broad corpus activity, unique objects accessed per day, requests, and GB delivered without our retaining individual requests. It cannot reliably tell whether every download was made by a human. CloudFront plus WAF provides stronger bot signals for traffic routed through it, but public direct S3 access remains outside that view.')
sub('Six proposals and their monthly costs')
table(['Proposal', 'What we learn', '80M', '800M', '8B'], [
    ['A  S3 native aggregates', 'Corpus totals, GB, distinct objects; bot status unknown', '$3–10', '$3–10', '$3–10'],
    ['B  CloudFront native reports', 'Popular objects, viewer totals, classified bot requests', '$203–210*', '$2,253–2,260*', 'Custom*'],
    ['C  Lambda log aggregation', 'Every file seen, GB, range counts, heuristic bot classes', '$5–20', '$50–100', '$500–1,000'],
    ['D  Athena daily aggregation', 'Same log detail with managed SQL and brief raw retention', '$5–15', '$20–60', '$160–350'],
    ['E  CloudTrail aggregation', 'API and resource access trends; no aggregate GB', '$110–130', '$1,050–1,100', '$10,450–10,650'],
    ['F  CloudFront 1 percent sample', 'Estimated file and bot patterns, range behavior', '$15–25†', '$15–30†', '$20–50†'],
], [1.65, 2.3, 1.05, 1.05, 1.05])
p('The user confirmed that existing S3 is free and CloudFront would be paid. Costs below therefore budget the new statistics services separately. Corpus storage and direct S3 delivery are $0 to the project under that arrangement. B includes a paid delivery plan; F excludes its separate paid delivery and WAF bill. Traffic columns are modeled workloads, not current monthly measurements. Labor, support plans, taxes, and DreamHost are excluded.')
p('* B assumes listed requests reach the distribution and transfer fits the plan. Business is $200 for 125M requests and 50 TB; $2,250 Premium provides 1.25B requests and 125 TB. At 8B requests, published levels stop at 6B, requiring a quote or reduced traffic. Existing S3 origin work is free under the stated arrangement. See page 5. [9,10]')
p('† F excludes CloudFront pay-as-you-go delivery and WAF, which can cost far more than sampling. Real-time sampling is incompatible with the flat-rate plans. See page 8. [10,20]')

page('Evidence and realistic sizing')
sub('What was verified for this report')
p('A complete anonymous public listing on 9 October found 19,266 current objects, 18,835,755,688,656 bytes (18.84 decimal TB), and six objects larger than 50 GiB. The bucket is publicly documented in us-west-2. Its twelve corpus categories include drives, files, mobile, packets, ram, scenarios, and several derived collections. This listing is not an inventory of hidden versions or a measurement of downloads. [1]')
p('The inspected checkout is digitalcorpora/digitalcorpora-stats, branch slg-dev, revision 7c249f9c62e6c7e529ef9b94b0fba638a687a6ef. It has unfinished DynamoDB and documentation changes. Its parser and design documents establish the S3 access-log → MySQL path; they do not prove the current deployed revision. Nothing was deployed or changed in AWS.')
sub('What the incident record establishes')
p('The earlier investigation recorded approximately 26.26M raw database rows for 21 September and 2.646M for 22 September. Traffic included repeated HTTP 206 ZIP probes using python-requests/2.32.5; one source was associated with roughly 3.36M rows and 2.82 TB. Partial ingestion retries also duplicated database rows. These figures describe the incident investigation, not fresh measurements or deduplicated source requests. They support burst planning, but do not establish that an AI company caused the traffic.')
table(['Sizing case', 'Requests per month', 'Average per day', 'Raw log assumption'], [
    ['Lower incident scale', '80 million', '2.67 million', '80 GB per month'],
    ['Peak scale sustained', '800 million', '26.7 million', '800 GB per month'],
    ['Tenfold stress', '8 billion', '267 million', '8,000 GB per month'],
], [1.9, 1.8, 1.55, 1.85])
p('Log calculations assume 1,000 uncompressed bytes per source request, a 30-day month, and no compression benefit. The raw log size, request mix, delivery-object size, and distinct-file fanout need measurement before implementation. Prices use binary GiB/TiB where AWS meters them; public GB/day would use decimal GB and label the units.')
sub('Sponsorship changes the decision')
p('The user confirmed during preparation that S3 is free to the project and CloudFront is not. The public site also documents sponsorship. AWS’s handbook permits temporary access logs with one-month expiry and expects analysis outside the sponsored account. Keep the free corpus route and budget Lambda, DynamoDB, monitoring, and other analysis in a paid account. Confirm whether new Storage Lens and metric fees are covered; this report conservatively budgets them as paid. [2,3]')
p('Live limitations: default and simsong AWS logins were expired. SSH rejected a host key that differs from etc/known_hosts. Private log volumes, enabled metrics, CloudFront configuration, and the deployed revision remain unverified. The public stats report could not be fetched. Sponsorship treatment follows the user’s explicit confirmation; private analytics usage remains unmeasured.', True)

page('What AWS can report without our request logs')
table(['AWS capability', 'Useful result', 'Main limit'], [
    ['Storage Lens Advanced', 'Daily corpus request and byte totals, distinct objects accessed, read-size histograms, errors', 'No per-object ranking or human identification'],
    ['Expanded prefix reports', 'Daily activity at small and deep corpus directory prefixes; CSV or Parquet export', 'Prefixes overlap; choose one reporting level'],
    ['S3 request metrics', 'Near real-time GETs, bytes, errors, latency by selected prefix or tag filters', 'GET count includes range requests and some internal operations'],
    ['CloudFront native reports', 'Top 50 objects, request and byte totals, incomplete responses, country and browser trends', 'Only CloudFront traffic; limited report history and approximations'],
    ['WAF bot metrics', 'Recognized bot names and categories, verified status, allowed and blocked counts', 'Request classifications, not completed files or GB by class'],
    ['CloudTrail aggregation', 'Five-minute API, resource, and principal counts', 'Requires paid atomic events; counts only in aggregate schema'],
    ['Billing and usage exports', 'Daily billable service quantities and costs', 'No file popularity or bot/human split'],
], [1.7, 3.15, 2.25])
p('Storage Lens is substantially more capable than older descriptions suggest. Its advanced tier includes performance metrics and expanded prefix activity reports at no additional feature fee. Export to ordinary S3 avoids the storage, maintenance, and query charges of S3 Tables. Prefix data is exported or viewed in S3 rather than published as prefix CloudWatch series. [4,5,6]')
p('The specific daily metric UniqueObjectsAccessedDailyCount is a useful “files accessed/day” measure. It records access, not successful full copies, and the same object accessed on two days counts once on each day. Adding daily unique counts cannot produce unique files for the month. [5]')
p('CloudFront generates popular-object and viewer reports without our enabling access logging. Its popular list covers the top 50 over the prior 60 days; low-ranked counts can be estimated. WAF labels and actions are available as metrics without our retaining full WAF logs. Free WAF bot visibility uses a sample and must be described as sampled. [7,8,11]')
p('S3 object-created notifications do not fire for GET downloads. They are useful for receiving a delivered log object or inventory changes, not as a free stream of download notifications. There is no S3 “file fully downloaded by a human” event. CloudWatch metric filters also require log ingestion first, so they do not eliminate the underlying collection. [12]')

page('Proposal A uses native S3 aggregates')
p('Best first choice when broad corpus popularity and volume are the priorities. Keep the existing public S3 download route. Enable one bucket-scoped Storage Lens Advanced dashboard with activity, performance, and expanded-prefix export. Publish daily corpus totals from the aggregate export, using one scheduled Lambda and small S3 JSON artifacts. DynamoDB is optional.')
sub('What we could learn')
p('Requests and bytes by corpus directory; unique objects accessed per day; read-size distributions that distinguish small probes from larger transfers; error trends; average read size; and activity relative to corpus size. This answers “which collections are being accessed?” and “how many GB and distinct files per day?” without building a request database. It leaves bot status explicitly unknown. [4,5]')
p('Use disjoint selected directories for public totals, with an “other” category. Nested prefixes must not be summed together. Do not use Storage Lens groups for activity grouping: advanced activity metrics are not available at the group level. Pilot the expanded prefix export before depending on any desired depth or leaf-level behavior. [5,6]')
sub('Monthly cost')
table(['Component', 'Planning quantity', 'Monthly cost'], [
    ['Storage Lens Advanced', '19,266 current objects × $0.20 per million', '$0.0039'],
    ['Aggregate export and public JSON', 'About 1 GB storage and 2,000 PUTs', 'About $0.04'],
    ['Scheduled Lambda', '30 runs × 60 seconds × 0.5 GB', 'About $0.02'],
    ['Dashboard and basic alarms', 'One paid dashboard and three alarms', 'About $3.30'],
    ['API reads, small state, diagnostics', 'Low volume; capped retention', 'Reserve $1–5'],
], [2.25, 3.4, 1.45])
p('Budget $3–10/month with paid dashboard and modest reporting. The pure metrics fee is under a cent at the observed object count; actual monitored count can include object versions. Removing the paid dashboard lowers the budget. Large expanded exports could increase the allowance, so select only necessary metrics and estimate export size in the pilot. Costs scale with object inventory and reporting activity rather than GET traffic. [4,13,14]')
sub('Optional faster monitoring')
p('If daily freshness is insufficient, add S3 request metrics for the twelve corpus categories plus a corpus-total and bucket-total filter. Fourteen filters with 7–16 active metrics at $0.30/metric-month imply roughly $29–67/month extra; a bucket-only configuration is roughly $2–5. Enabling a filter can emit and bill more metrics than the two shown in a dashboard. Do not create one filter per file. [13,15]')
p('Limits: both S3 request metrics and access logs are best-effort observations, not complete billing accounting. With a future CloudFront cache, S3 metrics measure origin work rather than bytes served to viewers; viewer reporting would need CloudFront. Use Storage Lens daily values as daily snapshots, replacing corrections rather than adding repeated exports. [15,16]')

page('Proposal B uses CloudFront and WAF reports')
p('Best native option for stronger bot visibility. Route web downloads through CloudFront, enable WAF Common Bot Control in count mode, and use native popular-object, viewer, and delivery reports. Persist only daily metrics and exported summaries. Standard access logs need not be enabled for popular-object and viewer reports. [7,8,11]')
sub('What we could learn')
p('Popular files with viewer bytes and request counts; GB/day including cache hits; viewer location and browser trends; request failures and incomplete-response indicators; recognized bot categories, bot names, verified status, and allowed versus blocked requests. WAF evaluates requests rather than downloaded response bodies, so these native metrics do not provide GB/day by bot class or a per-file bot cross-tab.')
table(['Flat rate tier', 'Monthly price', 'Request allowance', 'Transfer allowance'], [
    ['Pro', '$15', '10M', '50 TB'],
    ['Business with common bots', '$200', '125M', '50 TB'],
    ['Premium default', '$1,000', '500M', '50 TB'],
    ['Premium larger example', '$2,250', '1.25B', '125 TB'],
    ['Largest published Premium', '$10,000', '6B', '600 TB'],
], [2.25, 1.35, 1.75, 1.75])
p('At 80M requests, budget $203–210 including a small publisher if transfer fits 50 TB. Sustaining the historical one-source 2.82 TB/day would exceed that allowance; $1,450 Premium includes 75 TB and 750M requests and $2,250 includes 125 TB and 1.25B. At 800M, budget $2,253–2,260. At 8B, obtain a custom quote; $10,000 is not a sufficient published request allowance. S3 origin work remains free under the stated arrangement; optional logs and support are extra. Pro lacks full bot management. [9,10]')
p('Flat plans have no overage fees, but AWS may adjust delivery when sustained usage substantially exceeds the allowance. Treat the allowance as a sizing constraint. The plans include common bot control, not targeted bot control, and exclude real-time access logs. Do not add separate pay-as-you-go WAF charges to an included plan feature. [10]')
sub('What makes this conditional')
p('Coverage and expense: CloudFront is paid, while direct S3 is free. Moving delivery therefore adds cost rather than saving the project’s current S3 bill. Direct public S3 must remain available unless separately approved otherwise, so CloudFront sees a subset and bots can bypass it. Restricting the bucket would change existing access and may conflict with open-data distribution. This option is justified by better web-traffic bot visibility, not by delivery savings. [1,3]')
p('Large files: the public listing found six objects over 50 GiB. With caching enabled, oversized full-object GETs can fail; range GETs of smaller parts are supported, or use a cache-disabled behavior for large full responses. Test browser GET, Range, HEAD, CLI access, redirects, TLS, and hashes before moving links. [17]')

page('Proposal C aggregates S3 logs with Lambda')
p('Best choice if per-file observed activity is required while preserving direct S3 access. This processes every delivered request record briefly, then keeps only aggregates. It eliminates the per-request database, but does not eliminate AWS creation and delivery of raw access logs.')
p('Flow: log object delivered → SQS → bounded Lambda parser → aggregate deltas and idempotency state → daily S3 JSON. Run the paid analysis in us-west-2 outside the sponsored account, with explicit cross-account permissions. Raw inputs remain in the source account only until durable processing completes, under an approved short retention policy. [3,12,16]')
sub('What we could learn')
p('All observed file keys and corpus categories, daily payload bytes, successful full-response counts, HTTP 206 counts, failed responses, unique keys accessed, and broad declared-bot, likely automation, browser-like, and unknown classes. Summarize in memory or bounded chunks before writing. The source log has request IDs, object keys, response bytes, object size, and user agent; it does not provide enough information to prove range coverage across a client session. [18]')
sub('Worked monthly cost for broad aggregates')
table(['Component', '80M', '800M', '8B'], [
    ['Lambda at 1 GB and 2 s per log object', '$2.68', '$26.83', '$268.27'],
    ['DynamoDB 30 WRUs per log object', '$1.50', '$15.00', '$150.00'],
    ['SQS operations', '$0.10', '$0.96', '$9.60'],
    ['Existing sponsored S3 inputs', '$0', '$0', '$0'],
    ['Reporting, alerts, logs, small table', '$3–10', '$3–15', '$5–25'],
    ['Practical budget with retry headroom', '$5–20', '$50–100', '$500–1,000'],
], [3.55, 1.18, 1.18, 1.19])
p('Assumptions: 1,000 records per log object; N/1,000 invocations; three standard SQS operations per object; 30 transactional write units total for small aggregate deltas plus a marker; no secondary indexes, NAT gateway, or customer-managed key. Compute and write prices use $0.0000166667/GB-second and $0.625/million WRUs. Paid-service free tiers are excluded. Sponsored source S3 reads and storage cost the project $0; paid-account output storage is included in the reporting reserve. [4,14,19,22]')
p('Sensitivity: at 100 records/object, invocation, queue, and marker overhead can approach ten times the modeled amount. At 200 transactional WRUs/object, the write cost becomes $10/$100/$1,000 instead of $1.50/$15/$150. Exact file × classifier × status fanout can create this increase. Benchmark cold starts, parse time, batch density, retry rate, and fanout before accepting the budget.')
p('Avoid one DynamoDB write per request: even one ≤1 KB nontransactional write per request costs $50/$500/$5,000 in these cases. Transactional writes double that, before indexes and retries. BatchWriteItem does not make writes free or increment counters atomically. Store bounded precomputed aggregates, with retries made idempotent and popular counters sharded. [19]')

page('Proposals D and E use managed analysis')
sub('D uses Athena over short lived S3 logs')
p('Keep access logs briefly in S3, partition by event date, and run one scheduled Athena query per day to produce aggregate tables or JSON. Delete inputs after durable summarization under an approved retention policy. This provides file popularity, GB/day, response status, and user-agent heuristics with less custom parsing than C. It still captures request logs temporarily; it is not a no-log solution.')
p('At $5/TiB scanned, one pass over a 1,000-byte-per-request monthly input costs approximately $0.36/$3.64/$36.38 for 80M/800M/8B. Four passes to accommodate late arrivals cost $1.46/$14.55/$145.52. Existing sponsored source S3 adds $0. Budget $5–15/$20–60/$160–350 including paid-account results, orchestration, monitoring, and retry headroom. Compression and Parquet can lower scan charges, but conversion has its own cost. [4,21]')
p('Requirements: date-partitioned paths, a validated parser for quoted fields, and replacement of a day’s result when reprocessed. Thousands of tiny files can dominate query overhead and latency. Estimate S3 GET charges as input objects read × query passes × $0.0004/1,000; compact if justified. Use a scan limit per query and capped scheduled queries. Repeatedly scanning a growing month would invalidate these prices. [4,18,21]')
p('A useful variation for A is Athena over Storage Lens aggregate Parquet, rather than raw access logs. That retains the minimal-data model and costs much less to scan. Prefer ordinary S3 export first; S3 Tables adds managed storage and maintenance but no new bot evidence. [6]')
sub('E uses CloudTrail five minute aggregated events')
p('AWS offers real higher-level log events: CloudTrail can aggregate data events into five-minute API Activity, Resource Access, and User Actions summaries. This removes our need to parse individual audit events. However, enabling aggregation first requires a trail collecting data events, and the aggregation fee is additional to their capture fee. [23]')
p('What we learn: operation frequency, resource trends, errors, and client or identity breakdowns where included by the template. AWS’s S3 examples aggregate bucket ARNs; do not assume that RESOURCE_ACCESS supplies exact file rankings. The aggregate schema exposes counts, not response bytes, so GB/day still needs Storage Lens or S3 metrics. User-agent distributions support heuristics, not verified humanity. Marginal breakdowns cannot be treated as joint user × object counts. [23,24]')
p('CloudTrail retail cost is $1/million atomic data events plus $0.30/million analyzed for aggregation: $104/$1,040/$10,400 before storage and reporting. Budget $110–130/$1,050–1,100/$10,450–10,650, with actual atomic log storage size measured in a pilot. Additional copies or analysis add cost. Do not quote only the cheaper aggregation fee. [25]')
p('Assessment: useful if security auditing is also required, but a poor primary match for this project’s volume and GB questions. It still captures atomic events in AWS and is more expensive than S3 native metrics. Do not enable it solely to replace the current stats importer.')

page('Proposal F samples CloudFront requests')
p('Use AWS-configured 1 percent CloudFront real-time logging into a provisioned Kinesis stream, then a batched Lambda consumer that retains aggregates only. Keep native total request and byte metrics as the denominator. Select URI, method, status, bytes, user agent, and response range fields; omit cookies, query strings, referrer, and IP unless a specific analysis needs them. [20]')
sub('What we could learn and what we could not')
p('Estimate activity for high-volume files and corpus categories, range usage, declared bot or automation patterns, and changes in bytes per request. Extrapolate an observed sampled count by 100 only with an explicit sampling label. Rare files may disappear entirely. Do not multiply sampled unique-key counts by 100 to claim unique files/day. Best-effort delivery and clustered traffic further limit precision.')
p('For a 1 percent sample, a file with 100 requests/day has an approximately 36.6 percent probability of receiving no sample. A file with 1,000 requests/day yields about ten records: its approximate 95 percent count uncertainty is ±62 percent. With 10,000 requests/day, about 100 records yields ±20 percent. These binomial approximations assume representative independent sampling, which must be validated. Byte estimates can be much noisier when a few very large responses dominate.')
sub('Monthly mechanism cost')
p('Samples are 0.8M/8M/80M records. One provisioned Kinesis shard at $0.015/hour costs $10.80 per 30-day month; its PUT units add about $0.01/$0.11/$1.12 for small records. CloudFront real-time log fees at $0.01/million generated lines add $0.008/$0.08/$0.80. Batched Lambda, bounded aggregate writes, and alerts bring a practical budget to $15–25/$15–30/$20–50. Scale shards for peak arrivals, not monthly averages; 8B/month sampled averages about 31 records/s, but bursts may be much larger. [20,26,27]')
p('Critical restriction: real-time access logs are unsupported by CloudFront flat-rate plans. This requires pay-as-you-go CloudFront, with paid delivery and WAF separate. One source delivering 2.82 decimal TB/day sustained produces 84.6 TB/month: about $5,800 in US-edge retail transfer, plus HTTPS request charges before free allowances. Since existing S3 is free, this is an added expense. The byte figure is a historical extrapolation, not a current bill or full-corpus total. [10,27]')
sub('Bot cost and confidence')
p('On pay-as-you-go WAF, one ACL and Common Bot Control rule group cost about $16/month fixed, plus $0.60/million all WAF requests and $1/million inspected common-bot requests beyond 10M. Full inspection costs approximately $134/$1,286/$12,806 across our three scenarios, before additional rules or logs. Targeted Bot Control costs more and is not justified without evidence. [28]')
p('Keep named or verified bots, likely automation, browser-like, and unknown as separate classes. Researchers routinely use Python, curl, and range GETs. WAF bot labels improve attribution but do not prove human intent. Request-class percentages cannot be multiplied by total GB to estimate bot GB; that requires response-byte observations joined to classification. Native CloudFront samples do not include WAF labels automatically.')

page('Recommended replacement and acceptance criteria')
sub('Start with the information actually needed')
p('Choose A as the foundation: daily corpus and dataset prefix totals, distinct objects accessed/day, GB/day, read-size patterns, and error counts. Publish aggregate JSON through the existing site or a small static report. Store date, category, source, count definition, bytes, freshness, and a quality flag. Omit permanent IP histories and full user-agent strings unless a separately justified analysis needs them.')
p('Use a scheduled Lambda to ingest aggregate exports, not a Lambda per GET. S3 alone is sufficient for immutable daily snapshots and a current manifest. Add DynamoDB only for keyed daily data, processing state, or a small private reporting API. Keep inventory and cryptographic hashes as a separate responsibility; the new stats system need not reimplement hashing or persist old request-level dimensions.')
p('If per-file rankings are essential, add C or D with bounded, temporary inputs and only daily file aggregates. D has the lowest modeled managed-analysis cost if partitioning is sound; C gives more explicit recovery control. For bot identification on free S3, use temporary logs for honest heuristics. Stronger WAF classification requires the paid CloudFront option B for a web-facing subset. F’s low sample-processing cost does not offset its paid delivery bill.')
sub('Define the public numbers precisely')
table(['Published label', 'Definition and limit'], [
    ['Files accessed per day', 'Unique observed object keys in the day; not completed copies'],
    ['Download requests per day', 'Successful GET responses, with 200 and 206 separately where available'],
    ['GB delivered per day', 'Sum of observed bytes / 1,000,000,000; label origin versus viewer'],
    ['Full size responses', '200 responses with matching object size and bytes where logs allow; no client receipt proof'],
    ['Bot classification', 'Named or verified bots, likely automation, browser-like, unknown; label coverage'],
], [2.0, 5.1])
page('Checks before implementation and switchover')
p('Verify the renewed account session, coverage of new analytics fees, existing metric configuration, account permissions, export support, and the independently trusted server key. Measure seven consecutive days of deduplicated source requests, log bytes, records per delivery object, range fraction, and aggregate fanout. Confirm AWS-managed export delivery and cross-account reads through their real execution path.')
p('Pilot A alongside current reporting. Compare byte and request trends against native usage, mark late days provisional, and preserve the source timestamp and coverage. For C or D, inject interruption, retry, duplicate source records, late arrivals, and corrections; a log-object marker alone does not remove duplicate request IDs across separate delivered objects. Verify replay cannot double counters or erase newer summaries.')
p('Then approve retention and the switchover separately. Existing historical totals may contain replay inflation and should retain a quality note. Do not re-enable cron, delete the old database, block public S3, or change download URLs as part of preparing this report. Budget alerts at $5/$10/$25 are useful for A but require realistic revised thresholds if choosing a more expensive proposal.')
p('Planning effort, not a contractor quote: A roughly 2–5 engineer-days; B 5–10 plus sponsor coordination and download acceptance; C 10–20; D 5–10; E 3–5; F 8–15. Ongoing cloud prices omit this implementation and maintenance labor.')

sources = [
('DigitalCorpora AWS registry and bucket region','https://registry.opendata.aws/digitalcorpora/'),
('DigitalCorpora S3 information and sponsorship','https://digitalcorpora.org/about-digitalcorpora/s3-information/'),
('AWS Open Data onboarding handbook pages 4 and 15 to 16','https://assets.opendata.aws/aws-onboarding-handbook-for-data-providers-en-US.pdf'),
('Amazon S3 pricing and Oregon rate feed','https://aws.amazon.com/s3/pricing/'),
('Storage Lens metric glossary','https://docs.aws.amazon.com/AmazonS3/latest/userguide/storage_lens_metrics_glossary.html'),
('Storage Lens performance metrics and expanded prefix exports','https://aws.amazon.com/blogs/aws/amazon-s3-storage-lens-adds-performance-metrics-support-for-billions-of-prefixes-and-export-to-s3-tables/'),
('CloudFront popular objects reports','https://docs.aws.amazon.com/AmazonCloudFront/latest/DeveloperGuide/popular-objects-report.html'),
('CloudFront viewers reports','https://docs.aws.amazon.com/AmazonCloudFront/latest/DeveloperGuide/viewers-reports.html'),
('CloudFront pricing','https://aws.amazon.com/cloudfront/pricing/'),
('CloudFront flat rate plans and unsupported features','https://docs.aws.amazon.com/AmazonCloudFront/latest/DeveloperGuide/flat-rate-pricing-plan.html'),
('WAF metrics and free sampled bot visibility','https://docs.aws.amazon.com/waf/latest/developerguide/waf-metrics.html'),
('S3 event notification types','https://docs.aws.amazon.com/AmazonS3/latest/userguide/notification-how-to-event-types-and-destinations.html'),
('CloudWatch pricing','https://aws.amazon.com/cloudwatch/pricing/'),
('Lambda pricing','https://aws.amazon.com/lambda/pricing/'),
('S3 metrics and dimensions','https://docs.aws.amazon.com/AmazonS3/latest/userguide/metrics-dimensions.html'),
('S3 access logging and best effort collection','https://docs.aws.amazon.com/AmazonS3/latest/userguide/ServerLogs.html'),
('CloudFront range GETs and 50 GB caching boundary','https://docs.aws.amazon.com/AmazonCloudFront/latest/DeveloperGuide/RangeGETs.html'),
('S3 access log format and response byte fields','https://docs.aws.amazon.com/AmazonS3/latest/userguide/LogFormat.html'),
('DynamoDB pricing','https://aws.amazon.com/dynamodb/pricing/'),
('CloudFront real time log sampling and fields','https://docs.aws.amazon.com/AmazonCloudFront/latest/DeveloperGuide/real-time-logs.html'),
('Athena pricing','https://aws.amazon.com/athena/pricing/'),
('SQS pricing','https://aws.amazon.com/sqs/pricing/'),
('CloudTrail data event aggregation','https://docs.aws.amazon.com/awscloudtrail/latest/userguide/aggregating-data-events.html'),
('CloudTrail aggregated record schema','https://docs.aws.amazon.com/awscloudtrail/latest/userguide/cloudtrail-event-reference-aggregated-events.html'),
('CloudTrail pricing including aggregation surcharge','https://aws.amazon.com/cloudtrail/pricing/'),
('Kinesis Data Streams pricing','https://aws.amazon.com/kinesis/data-streams/pricing/'),
('CloudFront FAQs and real time logging price','https://aws.amazon.com/cloudfront/faqs/'),
('WAF and Common Bot Control pricing','https://aws.amazon.com/waf/pricing/'),
]
def link(para, label, url):
    rel = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'),rel)
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'174B77'); pr.append(color)
    r.append(pr); text = OxmlElement('w:t'); text.text=label; r.append(text); hl.append(r); para._p.append(hl)

page('Sources and price verification')
p('Research date is 9 October 2026. Numbered citations refer to the linked primary sources below. Calculations are planning models using current public retail rates, not a statement of the DigitalCorpora account’s effective bill.')
p('The official AmazonS3 Oregon offer feed was downloaded with publication date 28 September 2026. It lists Storage Lens at $0.20/million monitored objects-month, standard storage at $0.023/GiB-month in the first tier, GETs at $0.0004/1,000, and PUTs at $0.005/1,000. CloudFront’s offer feed was dated 3 October 2026; US HTTPS requests are $1/million and transfer tiers begin at $0.085/$0.080/$0.060 per GiB. Free allowances and volume discounts may reduce modeled pay-as-you-go totals.')
para=p('Official Oregon S3 price feed: '); link(para,'AmazonS3 us west 2 offer','https://pricing.us-east-1.amazonaws.com/offers/v1.0/aws/AmazonS3/current/us-west-2/index.json')
para=p('Official CloudFront price feed: '); link(para,'AmazonCloudFront offer','https://pricing.us-east-1.amazonaws.com/offers/v1.0/aws/AmazonCloudFront/current/index.json')
for i,(name,url) in enumerate(sources[:14],1):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(6)
    para.add_run(f'[{i}] '); link(para,name,url)
page('Additional sources and calculation notes')
for i,(name,url) in enumerate(sources[14:],15):
    para=doc.add_paragraph(); para.paragraph_format.space_after=Pt(6)
    para.add_run(f'[{i}] '); link(para,name,url)
sub('Reproducible calculations')
p('A: 19,266 / 1,000,000 × $0.20 = $0.0038532/month for one advanced dashboard. Object versions or broader dashboard scope change the quantity.')
p('C: for N requests and L = N/1,000 log objects, Lambda = L × (2 GB-seconds × $0.0000166667 + $0.20/1,000,000). DynamoDB = L × 30 × $0.625/1,000,000. SQS = L × 3 × $0.40/1,000,000. Existing sponsored S3 inputs are $0. If inputs moved to paid S3, seven-day raw storage would be N × 1,000 / 2³⁰ × 7/30 × $0.023.')
p('D: four scans = N × 1,000 / 2⁴⁰ × $5 × 4. For late data, re-run and replace the affected day; the four-pass model is an assumption to budget, not a completeness guarantee.')
p('E: N / 1,000,000 × ($1 + $0.30). F: sampled records = 0.01 × N. Common Bot Control with full inspection = $16 + $0.60 × N/1,000,000 + max(0, N/1,000,000 − 10) × $1.')
p('Delivery example: 2.82 × 10¹² bytes/day × 30 / 2³⁰ = about 78,790 GiB/month. US CloudFront tiers price the first 10,240 GiB at $0.085, the next 40,960 at $0.080, and the remaining approximately 27,590 at $0.060: about $5,800 before free allowances. Rounded estimates vary with the unit interpretation of the historical TB figure and geography.')
p('Internal evidence: repository source and design documents inspected read-only at the revision on page 2; incident investigation record last updated 7 October. Historical row counts were not treated as deduplicated requests. Public inventory was obtained through an anonymous, fully paginated S3 listing with no continuation token remaining. Private billing and log object density could not be inspected.')
doc.core_properties.title = 'DigitalCorpora statistics system options'
doc.core_properties.subject = 'AWS native metrics, bot visibility, and serverless cost comparison'
doc.core_properties.author = 'Codex'
doc.core_properties.keywords = 'DigitalCorpora, AWS, S3 Storage Lens, Lambda, DynamoDB'
path=OUT/'DigitalCorpora statistics system options 2026-10-09.docx'
for root in [doc._element, doc.styles._element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(path)
print(path)
