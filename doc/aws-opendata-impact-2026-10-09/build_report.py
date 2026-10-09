#!/usr/bin/env python3
# Build an editable sponsorship report from reviewed, private content.
# Pydantic validates the supplied paragraphs, tables, and source references.
# python-docx writes a compact memo with repeatable typography and page breaks.
# Tables repeat their header rows; rows stay together across page boundaries.
# Explicit URLs become clickable references without exposing mailbox exports.
# This module does not fetch evidence, infer metrics, or send the report.
# Keep content and generated artifacts outside the public source checkout.
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, Field, model_validator


class Table(BaseModel):
    headings: list[str]
    rows: list[list[str]]

    @model_validator(mode='after')
    def rectangular(self):
        if not self.headings or any(len(row) != len(self.headings) for row in self.rows):
            raise ValueError('Tables require headings and equal-length rows')
        return self


class Section(BaseModel):
    title: str
    paragraphs: list[str] = Field(default_factory=list)
    table: Table | None = None
    after: list[str] = Field(default_factory=list)
    page_break: bool = False


class Reference(BaseModel):
    label: str
    detail: str
    url: str = ''


class Report(BaseModel):
    title: str
    subtitle: str
    prepared_by: str
    date: str
    sections: list[Section]
    references: list[Reference] = Field(default_factory=list)


def hyperlink(paragraph, label: str, url: str):
    relationship = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link, run, properties = OxmlElement('w:hyperlink'), OxmlElement('w:r'), OxmlElement('w:rPr')
    link.set(qn('r:id'), relationship)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '245A81'); properties.append(color)
    run.append(properties)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def build(report: Report, output: Path):
    doc = Document()
    page = doc.sections[0]
    page.top_margin = page.bottom_margin = Inches(.65)
    page.left_margin = page.right_margin = Inches(.75)
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor.from_string('202B33')
    doc.core_properties.title = report.title
    doc.core_properties.author = report.prepared_by
    doc.add_paragraph(report.title, 'Title')
    doc.add_paragraph(report.subtitle)
    meta = doc.add_paragraph(f'Prepared by {report.prepared_by}  |  {report.date}')
    meta.runs[0].font.size = Pt(9)
    for section in report.sections:
        if section.page_break: doc.add_page_break()
        doc.add_paragraph(section.title, 'Heading 1')
        for text in section.paragraphs: doc.add_paragraph(text)
        if section.table:
            table = doc.add_table(rows=1, cols=len(section.table.headings))
            table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, section.table.headings):
                cell.text = text
                for run in cell.paragraphs[0].runs: run.bold = True
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in section.table.rows:
                row = table.add_row()
                for cell, text in zip(row.cells, values): cell.text = text
            for row in table.rows:
                keep = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(keep)
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs: run.font.size = Pt(9)
        for text in section.after: doc.add_paragraph(text)
    if report.references:
        doc.add_paragraph('Sources', 'Heading 1')
        for index, reference in enumerate(report.references, 1):
            paragraph = doc.add_paragraph(f'[{index}] ')
            if reference.url: hyperlink(paragraph, reference.label, reference.url)
            else: paragraph.add_run(reference.label)
            paragraph.add_run(f'. {reference.detail}')
            for run in paragraph.runs: run.font.size = Pt(9)
    footer = page.footer.paragraphs[0]
    footer.add_run('DigitalCorpora • AWS Open Data impact • ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(8)
    output.mkdir(parents=True, exist_ok=True)
    target = output / 'DigitalCorpora impact report for Kyle Cook.docx'
    doc.save(target)
    print(target)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--content', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(Report.model_validate_json(args.content.read_text()), args.output)


if __name__ == '__main__': main()
