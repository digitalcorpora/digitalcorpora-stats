# Verify evidence reconciliation and download counting with synthetic records.
# Exercise historical tally layouts, MIME/HTML extraction, and address identity.
# Confirm exclusions and explicit aliases change people counts without losing provenance.
# Execute the real shell scanner against repeated and failed synthetic S3 events.
# Validate document tables before writing an editable report and inspect its contents.
# All fixtures are invented and temporary; no real address or password is embedded.
# These tests do not contact Gmail, DreamHost, AWS, or the deployed responder.
import json
import stat
import subprocess
import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from pydantic import ValidationError

from build_report import Report, Section, Table, build
from request_audit import Decisions, Evidence, Headers, evidence_from_text, html_text, normalize_email, read_tally, summarize


class ReportToolsTest(unittest.TestCase):
    def test_tally_preserves_blank_columns_in_both_layouts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'tally.csv'
            for content in ['Date\tName\tEmail\tdomestic\tundergraduates\tgradautes\n2023-01-24\nPerson\np@example.edu\nno\n\n12\n',
                            'Date\tName\tEmail\tdomestic\tundergraduates\tgradautes\n2023-01-24\tPerson\tp@example.edu\tno\t\t12\n']:
                path.write_text(content)
                record, = read_tally(path)
                self.assertEqual((record.date, record.email, record.undergraduates, record.graduates), (date(2023, 1, 24), 'p@example.edu', '', '12'))

    def test_html_and_bom_form_select_first_quoted_request(self):
        text = html_text('<style>hidden</style><div>\ufeffName: Person<br>Email: P@example.edu<br>gradautes: 12</div><div>Name: Other<br>Email: other@example.edu</div>')
        record = evidence_from_text(text, Headers(date='Fri, 9 Oct 2026 10:00:00 +0000'), 'archive', '1')
        self.assertEqual((record.email, record.name, record.graduates), ('p@example.edu', 'Person', '12'))
        self.assertNotIn('hidden', text)

    def test_alias_exclusion_and_annual_first_observation(self):
        records = [Evidence(source='test', source_id='1', email='old@example.edu', date=date(2021, 1, 1), form=True),
                   Evidence(source='test', source_id='2', email='new@example.edu', date=date(2022, 1, 1), form=True),
                   Evidence(source='test', source_id='3', email='spam@example.com', date=date(2022, 1, 1), form=True),
                   Evidence(source='test', source_id='4', email='manual@example.edu', date=date(2022, 1, 1))]
        result = summarize(records, Decisions(aliases={'old@example.edu': 'new@example.edu'}, excluded={'spam@example.com': 'spam'}, confirmed_manual=['manual@example.edu']))
        self.assertEqual((result.records, result.confirmed_requesters, result.excluded_identities, result.unresolved_identities), (4, 2, 1, 0))
        self.assertEqual([year.first_observed_requesters for year in result.annual], [1, 1])
        self.assertEqual(normalize_email('First.Last+tag@googlemail.com'), 'firstlast@gmail.com')

    def test_downloads_deduplicate_ids_and_filter_errors(self):
        def line(identifier, operation='REST.GET.OBJECT', status='200', size='12', key='corpora/scenarios/2009-m57-patents/file.e01', year=2025):
            return f'owner bucket [01/Jan/{year}:00:00:00 +0000] 192.0.2.1 - {identifier} {operation} {key} "GET /x HTTP/1.1" {status} - {size} 12 1 1 "-" "agent"\n'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'access.log'
            path.write_text(line('A1ZY') * 2 + line('B2', status='206', size='3') + line('C3', status='404') + line('D4', operation='REST.HEAD.OBJECT') + line('E5', size='-', key='corpora/empty') + line('F6', key='stats/digitalcorpora_configuration1.csv'))
            completed = subprocess.run(['bash', str(Path(__file__).with_name('count_access_logs.sh')), str(path)], capture_output=True, text=True, check=True)
            self.assertIn('all\t3\t15', completed.stdout)
            self.assertIn('m57\t2\t15', completed.stdout)
            self.assertIn('malformed_successful_lines\t0', completed.stderr)
            other = Path(directory) / 'next.log'
            other.write_text(line('A1ZY') + line('G7', year=2026))
            scanner = str(Path(__file__).with_name('count_access_logs.sh'))
            direct = subprocess.run(['bash', scanner, str(path), str(other)], capture_output=True, text=True, check=True)
            self.assertIn('2025\tall\t3\t15', direct.stdout)
            self.assertIn('2026\tall\t1\t12', direct.stdout)
            parts = []
            for index, source in enumerate([path, other]):
                part = Path(directory) / f'part-{index}'
                with part.open('w') as output:
                    subprocess.run(['bash', scanner, '--extract', str(source)], stdout=output, stderr=subprocess.PIPE, check=True)
                parts.append(str(part))
            merged = subprocess.run(['bash', scanner, '--merge', *parts], capture_output=True, text=True, check=True)
            self.assertEqual(merged.stdout, direct.stdout)

    def test_document_has_reviewed_content_and_rejects_ragged_table(self):
        with self.assertRaises(ValidationError): Table(headings=['a', 'b'], rows=[['only one']])
        with tempfile.TemporaryDirectory() as directory:
            report = Report(title='Impact', subtitle='For recipient', prepared_by='Author', date='2026-10-09', sections=[Section(title='Counts', paragraphs=['Reviewed text'], table=Table(headings=['Year', 'Requests'], rows=[['2025', '3']]))])
            build(report, Path(directory))
            document = Document(Path(directory) / 'DigitalCorpora impact report for Kyle Cook.docx')
            self.assertIn('Reviewed text', [paragraph.text for paragraph in document.paragraphs])
            self.assertEqual(document.tables[0].cell(1, 1).text, '3')

    def test_audit_enforces_private_permissions(self):
        import sys
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'source', Path(directory) / 'audit'
            source.mkdir(); output.mkdir(mode=0o755)
            (source / 'archive.mbox').write_text('')
            for name in ('tally-all.csv', 'tally.csv', 't2.csv'):
                (source / name).write_text('Date\tName\tEmail\n')
            (output / 'records.json').write_text('old data')
            (output / 'records.json').chmod(0o644)
            subprocess.run([sys.executable, str(Path(__file__).with_name('request_audit.py')), '--input', str(source), '--output', str(output)], capture_output=True, check=True)
            self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o700)
            for name in ('records.json', 'summary.json'):
                self.assertEqual(stat.S_IMODE((output / name).stat().st_mode), 0o600)


if __name__ == '__main__': unittest.main()
